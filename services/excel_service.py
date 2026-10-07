import io
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_file(records: List[Dict[str, Any]], schema_fields: List[Dict[str, str]]) -> io.BytesIO:
    """
    Tạo file Excel tổng hợp các tài liệu đã trích xuất với định dạng chuẩn chuyên nghiệp.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Dữ liệu Trích Xuất"
    ws.views.sheetView[0].showGridLines = True

    # Định dạng Font & Màu sắc
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Segoe UI", size=10)
    data_align = Alignment(vertical="center", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    # 1. Hàng tiêu đề chính
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "BÁO CÁO TỔNG HỢP DỮ LIỆU TÀI LIỆU (DOCUBATCH AI)"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30

    # 2. Hàng Header cột
    headers = ["STT", "Tên tệp gốc"] + [f.get("label", f.get("id")) for f in schema_fields] + ["Tóm tắt AI", "Trạng thái đối soát"]
    ws.append([]) # Dòng 2 để trống
    ws.append(headers) # Dòng 3 là Header
    ws.row_dimensions[3].height = 28

    header_row_idx = 3
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # 3. Dữ liệu các dòng
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    for i, item in enumerate(records, start=1):
        row_num = header_row_idx + i
        ws.row_dimensions[row_num].height = 24

        filename = item.get("filename", "")
        extracted = item.get("extractedData", {})
        summary = extracted.get("_ai_summary", "")
        verified = "Đã đối soát" if item.get("verified", False) else "Chưa đối soát"

        row_values = [i, filename]
        for field in schema_fields:
            fid = field.get("id")
            val = extracted.get(fid, "")
            row_values.append(str(val) if val is not None else "")

        row_values.append(summary)
        row_values.append(verified)

        ws.append(row_values)

        # Style từng ô trong hàng
        use_zebra = (i % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if use_zebra:
                cell.fill = zebra_fill

            # Căn giữa cho STT và Trạng thái
            if col_idx in (1, len(headers)):
                cell.alignment = center_align
            else:
                cell.alignment = data_align

    # 4. Tự động căn chỉnh độ rộng cột
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 3:
                continue
            if cell.value:
                # Ước lượng chiều dài
                length = len(str(cell.value))
                if length > max_len:
                    max_len = length
        ws.column_dimensions[col_letter].width = max(min(max_len + 5, 50), 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
